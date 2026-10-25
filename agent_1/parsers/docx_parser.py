"""DOCX parser using python-docx with paragraph level mapping."""

from pathlib import Path
from typing import Optional, List
import docx

from agent_1.parsers.base import AbstractDocumentParser
from agent_1.models.document import (
    ParsedDocument,
    DocumentParagraph,
    DocumentFormat,
)
from agent_1.models.manifest import ManifestEntry


class DocxParser(AbstractDocumentParser):
    """Parses Microsoft Word (.docx) documents into structured paragraphs."""

    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.DOCX

    def parse(
        self,
        file_path: Path,
        doc_id: str,
        manifest_meta: Optional[ManifestEntry] = None,
    ) -> ParsedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        doc = docx.Document(str(file_path))
        paragraphs: List[DocumentParagraph] = []
        full_text_chunks: List[str] = []
        para_index = 1

        # Process document body paragraphs
        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            cleaned = " ".join([line.strip() for line in text.splitlines() if line.strip()])
            paragraphs.append(
                DocumentParagraph.create(
                    index=para_index,
                    text=cleaned,
                    page_number=None,
                )
            )
            full_text_chunks.append(cleaned)
            para_index += 1

        # Process tables if present
        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    row_text = " | ".join(row_cells)
                    paragraphs.append(
                        DocumentParagraph.create(
                            index=para_index,
                            text=f"[Table Row]: {row_text}",
                            page_number=None,
                        )
                    )
                    full_text_chunks.append(row_text)
                    para_index += 1

        full_text = "\n\n".join(full_text_chunks)
        return ParsedDocument(
            doc_id=doc_id,
            file_name=file_path.name,
            file_path=str(file_path.resolve()),
            format=DocumentFormat.DOCX,
            manifest_meta=manifest_meta,
            paragraphs=paragraphs,
            full_text=full_text,
            total_paragraphs=len(paragraphs),
            total_pages=None,
        )
