"""DOCX parser using python-docx with standard library zipfile/xml fallback."""

from pathlib import Path
from typing import Optional, List
import logging

try:
    import docx
    HAS_DOCX = True
except ImportError:
    HAS_DOCX = False
    import zipfile
    import xml.etree.ElementTree as ET

from agent_1.parsers.base import AbstractDocumentParser
from agent_1.models.document import (
    ParsedDocument,
    DocumentParagraph,
    DocumentFormat,
)
from agent_1.models.manifest import ManifestEntry

logger = logging.getLogger(__name__)


class DocxParser(AbstractDocumentParser):
    """Parses Microsoft Word (.docx) documents into structured paragraphs."""

    @property
    def supported_format(self) -> DocumentFormat:
        return DocumentFormat.DOCX

    def parse(
        self,
        file_path: Path,
        doc_id: str,
        manifest_meta: Optional[ManifestEntry] = None,
    ) -> ParsedDocument:
        if not file_path.exists():
            raise FileNotFoundError(f"DOCX file not found: {file_path}")

        if HAS_DOCX:
            return self._parse_with_docx(file_path, doc_id, manifest_meta)
        else:
            return self._parse_with_zipfile(file_path, doc_id, manifest_meta)

    def _parse_with_docx(
        self,
        file_path: Path,
        doc_id: str,
        manifest_meta: Optional[ManifestEntry] = None,
    ) -> ParsedDocument:
        doc = docx.Document(str(file_path))
        paragraphs: List[DocumentParagraph] = []
        full_text_chunks: List[str] = []
        para_index = 1

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue
            cleaned = " ".join([line.strip() for line in text.splitlines() if line.strip()])
            paragraphs.append(
                DocumentParagraph.create(
                    index=para_index,
                    text=cleaned,
                    page_number=None,
                )
            )
            full_text_chunks.append(cleaned)
            para_index += 1

        for table in doc.tables:
            for row in table.rows:
                row_cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if row_cells:
                    row_text = " | ".join(row_cells)
                    paragraphs.append(
                        DocumentParagraph.create(
                            index=para_index,
                            text=row_text,
                            page_number=None,
                        )
                    )
                    full_text_chunks.append(row_text)
                    para_index += 1

        full_text = "\n\n".join(full_text_chunks)
        return ParsedDocument(
            doc_id=doc_id,
            file_name=file_path.name,
            file_path=str(file_path.resolve()),
            format=DocumentFormat.DOCX,
            manifest_meta=manifest_meta,
            paragraphs=paragraphs,
            full_text=full_text,
            total_paragraphs=len(paragraphs),
            total_pages=None,
        )

    def _parse_with_zipfile(
        self,
        file_path: Path,
        doc_id: str,
        manifest_meta: Optional[ManifestEntry] = None,
    ) -> ParsedDocument:
        import zipfile
        import xml.etree.ElementTree as ET

        paragraphs: List[DocumentParagraph] = []
        full_text_chunks: List[str] = []
        para_index = 1

        with zipfile.ZipFile(str(file_path)) as docx_zip:
            xml_content = docx_zip.read('word/document.xml')
            tree = ET.fromstring(xml_content)
            # w:p represents paragraphs, w:t represents text runs
            namespaces = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
            for p in tree.iterfind('.//w:p', namespaces):
                texts = [node.text for node in p.iterfind('.//w:t', namespaces) if node.text]
                if texts:
                    cleaned = " ".join("".join(texts).split())
                    if cleaned:
                        paragraphs.append(
                            DocumentParagraph.create(
                                index=para_index,
                                text=cleaned,
                                page_number=None,
                            )
                        )
                        full_text_chunks.append(cleaned)
                        para_index += 1

        full_text = "\n\n".join(full_text_chunks)
        return ParsedDocument(
            doc_id=doc_id,
            file_name=file_path.name,
            file_path=str(file_path.resolve()),
            format=DocumentFormat.DOCX,
            manifest_meta=manifest_meta,
            paragraphs=paragraphs,
            full_text=full_text,
            total_paragraphs=len(paragraphs),
            total_pages=None,
        )
